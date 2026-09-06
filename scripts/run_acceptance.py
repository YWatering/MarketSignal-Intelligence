#!/usr/bin/env python3
"""Run deterministic end-to-end acceptance checks."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from market_batch import run_batch
from versioning import BATCH_CONTRACT_VERSION


def main() -> int:
    with tempfile.TemporaryDirectory() as temporary_directory:
        root = Path(temporary_directory)
        manifest = root / "acceptance.yaml"
        manifest.write_text(
            "\n".join(
                [
                    f'version: "{BATCH_CONTRACT_VERSION}"',
                    "name: fixture-acceptance",
                    "task_type: portfolio",
                    "mode: fixture",
                    "start_date: 2026-08-03",
                    "end_date: 2026-08-14",
                    f'output_dir: "{root / "items"}"',
                    f'summary_output: "{root / "summary.xlsx"}"',
                    f'log_file: "{root / "run.jsonl"}"',
                    f'state_file: "{root / "state.json"}"',
                    "items:",
                    "  - symbol: AAPL",
                    "    market: us",
                    "    label: Apple fixture",
                ]
            ),
            encoding="utf-8",
        )
        summary = run_batch(manifest, force=True)
        if summary["status"] != "pass":
            raise RuntimeError(f"acceptance batch failed: {summary}")
        workbook = load_workbook(summary["summary_output"], read_only=True, data_only=True)
        expected = ["README", "主体任务", "横向比较", "版本信息"]
        if workbook.sheetnames != expected:
            raise RuntimeError(f"unexpected batch workbook sheets: {workbook.sheetnames}")
        workbook.close()
        print(
            json.dumps(
                {
                    "status": "pass",
                    "batch_status": summary["status"],
                    "successful_items": summary["successful_items"],
                    "workbook_sheets": expected,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
