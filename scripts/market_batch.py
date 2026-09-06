#!/usr/bin/env python3
"""Run multi-stock, industry, or theme tasks from a versioned manifest."""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from datetime import date
from pathlib import Path
from typing import Any, Callable

import yaml
from openpyxl import Workbook

from marketsignal import (
    PipelineError,
    _remove_core_properties,
    _set_number_format,
    _write_table,
    parse_symbol,
    run_pipeline,
)
from operations import JsonlRunLogger, retry_operation, schedule_is_due
from versioning import BATCH_CONTRACT_VERSION, component_versions


ALLOWED_TASK_TYPES = {"portfolio", "industry", "theme"}
ALLOWED_MODES = {"fixture", "online"}
COMPLETED_STATUSES = {"pass", "warning"}
PROJECT_ROOT = Path(__file__).resolve().parents[1]


class BatchError(RuntimeError):
    """Expected batch manifest or execution failure."""


def _read_manifest(path: Path) -> dict[str, Any]:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as exc:
        raise BatchError(f"could not read batch manifest: {path}") from exc
    if not isinstance(payload, dict):
        raise BatchError("batch manifest must contain a mapping at the top level")
    return payload


def _safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return cleaned or "item"


def _display_path(path: Path) -> str:
    try:
        return str(path.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def _as_optional_date(value: Any, field: str) -> str | None:
    if value in {None, ""}:
        return None
    text = str(value).strip()
    if text == "latest":
        return date.today().isoformat()
    try:
        date.fromisoformat(text)
    except ValueError as exc:
        raise BatchError(f"{field} must use YYYY-MM-DD, latest, or be omitted") from exc
    return text


def _number_setting(value: Any, field: str, number_type: type[int] | type[float]) -> int | float:
    try:
        return number_type(value)
    except (TypeError, ValueError) as exc:
        raise BatchError(f"{field} must be a valid {number_type.__name__}") from exc


def validate_manifest(payload: dict[str, Any]) -> dict[str, Any]:
    version = str(payload.get("version", ""))
    if version != BATCH_CONTRACT_VERSION:
        raise BatchError(
            f"batch manifest version must be {BATCH_CONTRACT_VERSION}; received {version or 'missing'}"
        )
    name = str(payload.get("name", "")).strip()
    if not name:
        raise BatchError("batch manifest requires a non-empty name")
    task_type = str(payload.get("task_type", "")).strip()
    if task_type not in ALLOWED_TASK_TYPES:
        raise BatchError(f"task_type must be one of: {', '.join(sorted(ALLOWED_TASK_TYPES))}")
    mode = str(payload.get("mode", "online")).strip()
    if mode not in ALLOWED_MODES:
        raise BatchError("mode must be fixture or online")
    items = payload.get("items")
    if not isinstance(items, list) or not items:
        raise BatchError("batch manifest requires at least one item")
    if len(items) > 50:
        raise BatchError("batch manifest supports at most 50 items per run")
    normalized_items = []
    identities: set[tuple[str, str]] = set()
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            raise BatchError(f"item {index} must be a mapping")
        symbol = str(item.get("symbol", "")).strip()
        market = str(item.get("market", "")).strip() or None
        try:
            instrument = parse_symbol(symbol, market)
        except PipelineError as exc:
            raise BatchError(f"item {index}: {exc}") from exc
        identity = (instrument["symbol"], instrument["market"])
        if identity in identities:
            raise BatchError(f"duplicate batch item: {instrument['symbol']}")
        identities.add(identity)
        normalized_items.append(
            {
                "symbol": symbol,
                "normalized_symbol": instrument["symbol"],
                "market": instrument["market"],
                "label": str(item.get("label", instrument["symbol"])).strip() or instrument["symbol"],
                "group": str(item.get("group", payload.get("group", ""))).strip(),
            }
        )
    output_dir = Path(str(payload.get("output_dir", "outputs/batch")))
    summary_output = Path(str(payload.get("summary_output", output_dir / "summary.xlsx")))
    if summary_output.suffix.lower() != ".xlsx":
        raise BatchError("summary_output must use the .xlsx extension")
    execution = payload.get("execution") or {}
    if not isinstance(execution, dict):
        raise BatchError("execution must be a mapping")
    attempts = int(_number_setting(execution.get("attempts", 2), "execution.attempts", int))
    if attempts < 1 or attempts > 5:
        raise BatchError("execution.attempts must be between 1 and 5")
    retry_delay_seconds = float(
        _number_setting(execution.get("retry_delay_seconds", 1.0), "execution.retry_delay_seconds", float)
    )
    if retry_delay_seconds < 0 or retry_delay_seconds > 60:
        raise BatchError("execution.retry_delay_seconds must be between 0 and 60")
    schedule = payload.get("schedule") or {}
    if not isinstance(schedule, dict):
        raise BatchError("schedule must be a mapping")
    interval_hours = schedule.get("interval_hours")
    if interval_hours is not None:
        interval_hours = float(_number_setting(interval_hours, "schedule.interval_hours", float))
        if interval_hours < 1:
            raise BatchError("schedule.interval_hours must be at least 1")
    forecast = bool(payload.get("forecast", False))
    if forecast and mode != "online":
        raise BatchError("formal batch forecasts require online mode")
    start_date = _as_optional_date(payload.get("start_date"), "start_date")
    end_date = _as_optional_date(payload.get("end_date"), "end_date")
    if start_date and end_date and start_date > end_date:
        raise BatchError("start_date cannot be later than end_date")
    forecast_horizon = int(_number_setting(payload.get("forecast_horizon", 5), "forecast_horizon", int))
    forecast_minimum_history = int(
        _number_setting(payload.get("forecast_minimum_history", 120), "forecast_minimum_history", int)
    )
    forecast_validation_points = int(
        _number_setting(payload.get("forecast_validation_points", 40), "forecast_validation_points", int)
    )
    forecast_ridge_alpha = float(
        _number_setting(payload.get("forecast_ridge_alpha", 1.0), "forecast_ridge_alpha", float)
    )
    cross_validation_tolerance_pct = float(
        _number_setting(
            payload.get("cross_validation_tolerance_pct", 1.0),
            "cross_validation_tolerance_pct",
            float,
        )
    )
    if forecast_horizon < 1 or forecast_horizon > 20:
        raise BatchError("forecast_horizon must be between 1 and 20")
    if forecast_minimum_history < 80:
        raise BatchError("forecast_minimum_history must be at least 80")
    if forecast_validation_points < 10:
        raise BatchError("forecast_validation_points must be at least 10")
    if forecast_ridge_alpha <= 0:
        raise BatchError("forecast_ridge_alpha must be greater than zero")
    if cross_validation_tolerance_pct <= 0:
        raise BatchError("cross_validation_tolerance_pct must be greater than zero")
    return {
        **payload,
        "name": name,
        "task_type": task_type,
        "mode": mode,
        "items": normalized_items,
        "start_date": start_date,
        "end_date": end_date,
        "output_dir": output_dir,
        "summary_output": summary_output,
        "log_file": Path(str(payload.get("log_file", ".cache/marketsignal/batch.jsonl"))),
        "state_file": Path(str(payload.get("state_file", ".cache/marketsignal/batch-state.json"))),
        "execution": {
            "attempts": attempts,
            "retry_delay_seconds": retry_delay_seconds,
            "continue_on_error": bool(execution.get("continue_on_error", True)),
        },
        "schedule": {"interval_hours": interval_hours},
        "forecast": forecast,
        "forecast_horizon": forecast_horizon,
        "forecast_minimum_history": forecast_minimum_history,
        "forecast_validation_points": forecast_validation_points,
        "forecast_ridge_alpha": forecast_ridge_alpha,
        "cross_validate_prices": bool(payload.get("cross_validate_prices", mode == "online")),
        "cross_validation_tolerance_pct": cross_validation_tolerance_pct,
    }


def _read_state(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def _write_state(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _resolve_project_paths(config: dict[str, Any]) -> None:
    for field in ("output_dir", "summary_output", "log_file", "state_file"):
        path = config[field]
        if not path.is_absolute():
            config[field] = PROJECT_ROOT / path


def _build_batch_workbook(config: dict[str, Any], results: list[dict[str, Any]], status: str) -> None:
    output = config["summary_output"]
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    readme = workbook.create_sheet("README")
    _write_table(
        readme,
        ["field", "value", "notes"],
        [
            ["task_name", config["name"], "Versioned batch manifest task"],
            ["task_type", config["task_type"], "portfolio, industry, or theme"],
            ["group_definition", "user supplied constituents", "No automatic constituent guessing"],
            ["mode", config["mode"], "Data collection mode"],
            ["status", status, "pass, warning, partial, failed, or skipped"],
            ["requested_items", len(config["items"]), "Configured subjects"],
            ["successful_items", sum(row["status"] in COMPLETED_STATUSES for row in results), "Completed subjects"],
            ["warning_items", sum(row["status"] == "warning" for row in results), "Completed subjects with quality warnings"],
            ["failed_items", sum(row["status"] == "failed" for row in results), "Isolated failures"],
            ["batch_contract_version", BATCH_CONTRACT_VERSION, "Manifest schema version"],
        ],
    )
    item_sheet = workbook.create_sheet("主体任务")
    item_headers = ["label", "group", "symbol", "market", "status", "attempts", "output", "error"]
    _write_table(item_sheet, item_headers, [[row.get(key, "") for key in item_headers] for row in results])
    comparison_sheet = workbook.create_sheet("横向比较")
    comparison_headers = [
        "label",
        "group",
        "symbol",
        "market",
        "currency",
        "latest_price_date",
        "latest_close",
        "period_return_pct",
        "close_return_volatility_pct",
        "net_profit_margin_pct",
        "current_ratio",
        "operating_cash_flow_margin_pct",
        "positive_news_count",
        "negative_news_count",
        "news_tone_balance",
        "price_rows",
        "financial_rows",
        "news_rows",
        "announcement_rows",
        "forecast_status",
        "selected_model",
        "selected_forecast_close",
        "forecast_lower_bound",
        "forecast_upper_bound",
        "cross_validation_status",
        "status",
    ]
    comparison_rows = []
    for result in results:
        summary = result.get("summary", {})
        comparison_rows.append(
            [
                result["label"],
                result["group"],
                result["symbol"],
                result["market"],
                summary.get("currency", ""),
                summary.get("latest_price_date", ""),
                summary.get("latest_close", ""),
                summary.get("period_return_pct", ""),
                summary.get("close_return_volatility_pct", ""),
                summary.get("net_profit_margin_pct", ""),
                summary.get("current_ratio", ""),
                summary.get("operating_cash_flow_margin_pct", ""),
                summary.get("positive_news_count", 0),
                summary.get("negative_news_count", 0),
                summary.get("news_tone_balance", 0),
                summary.get("price_rows", 0),
                summary.get("financial_rows", 0),
                summary.get("news_rows", 0),
                summary.get("announcement_rows", 0),
                summary.get("forecast_status", ""),
                summary.get("selected_model", ""),
                summary.get("selected_forecast_close", ""),
                summary.get("forecast_lower_bound", ""),
                summary.get("forecast_upper_bound", ""),
                summary.get("cross_validation_status", "not_requested"),
                result["status"],
            ]
        )
    _write_table(comparison_sheet, comparison_headers, comparison_rows)
    _set_number_format(
        comparison_sheet,
        {
            "latest_close",
            "period_return_pct",
            "close_return_volatility_pct",
            "net_profit_margin_pct",
            "current_ratio",
            "operating_cash_flow_margin_pct",
            "selected_forecast_close",
            "forecast_lower_bound",
            "forecast_upper_bound",
        },
        "0.0000",
    )
    version_sheet = workbook.create_sheet("版本信息")
    _write_table(
        version_sheet,
        ["component", "version", "notes"],
        [[component, version, "central version registry"] for component, version in component_versions().items()],
    )
    workbook.save(output)
    _remove_core_properties(output)


def run_batch(
    manifest_path: Path,
    *,
    force: bool = False,
    dry_run: bool = False,
    runner: Callable[..., dict[str, Any]] = run_pipeline,
    now_epoch: float | None = None,
) -> dict[str, Any]:
    config = validate_manifest(_read_manifest(manifest_path))
    _resolve_project_paths(config)
    now_epoch = time.time() if now_epoch is None else now_epoch
    state = _read_state(config["state_file"])
    due = schedule_is_due(
        interval_hours=config["schedule"]["interval_hours"],
        last_completion_epoch=state.get("last_completion_epoch", state.get("last_success_epoch")),
        now_epoch=now_epoch,
    )
    if dry_run:
        return {
            "status": "validated",
            "name": config["name"],
            "task_type": config["task_type"],
            "items": len(config["items"]),
            "due": due,
        }
    if not due and not force:
        return {
            "status": "skipped",
            "name": config["name"],
            "reason": "configured interval has not elapsed",
            "summary_output": _display_path(config["summary_output"]),
        }
    logger = JsonlRunLogger(config["log_file"])
    logger.emit(
        "batch_started",
        task_name=config["name"],
        task_type=config["task_type"],
        item_count=len(config["items"]),
        versions=component_versions(),
    )
    results: list[dict[str, Any]] = []
    config["output_dir"].mkdir(parents=True, exist_ok=True)
    for item in config["items"]:
        output = config["output_dir"] / f"{_safe_name(item['normalized_symbol'].replace('.', '_'))}.xlsx"
        context = {
            "task_name": config["name"],
            "symbol": item["normalized_symbol"],
            "market": item["market"],
        }

        def execute_item() -> dict[str, Any]:
            return runner(
                symbol=item["symbol"],
                market=item["market"],
                start_date=config["start_date"],
                end_date=config["end_date"],
                mode=config["mode"],
                output=output,
                forecast=config["forecast"],
                forecast_horizon=config["forecast_horizon"],
                forecast_minimum_history=config["forecast_minimum_history"],
                forecast_validation_points=config["forecast_validation_points"],
                forecast_ridge_alpha=config["forecast_ridge_alpha"],
                cross_validate_prices=config["cross_validate_prices"],
                cross_validation_tolerance_pct=config["cross_validation_tolerance_pct"],
            )

        try:
            summary, attempts = retry_operation(
                execute_item,
                attempts=config["execution"]["attempts"],
                delay_seconds=config["execution"]["retry_delay_seconds"],
                logger=logger,
                context=context,
            )
            item_status = "warning" if summary.get("status") == "warning" else "pass"
            results.append(
                {
                    "label": item["label"],
                    "group": item["group"],
                    "symbol": item["normalized_symbol"],
                    "market": item["market"],
                    "status": item_status,
                    "attempts": attempts,
                    "output": _display_path(output),
                    "error": "",
                    "summary": summary,
                }
            )
        except Exception as exc:
            results.append(
                {
                    "label": item["label"],
                    "group": item["group"],
                    "symbol": item["normalized_symbol"],
                    "market": item["market"],
                    "status": "failed",
                    "attempts": config["execution"]["attempts"],
                    "output": _display_path(output),
                    "error": str(exc),
                    "summary": {},
                }
            )
            if not config["execution"]["continue_on_error"]:
                break
    successful = sum(result["status"] in COMPLETED_STATUSES for result in results)
    warnings = sum(result["status"] == "warning" for result in results)
    if successful == len(config["items"]) and warnings:
        status = "warning"
    elif successful == len(config["items"]):
        status = "pass"
    elif successful:
        status = "partial"
    else:
        status = "failed"
    _build_batch_workbook(config, results, status)
    if status in {"pass", "warning"}:
        _write_state(
            config["state_file"],
            {
                "task_name": config["name"],
                "last_completion_epoch": now_epoch,
                "batch_contract_version": BATCH_CONTRACT_VERSION,
            },
        )
    logger.emit(
        "batch_finished",
        task_name=config["name"],
        status=status,
        successful_items=successful,
        failed_items=sum(result["status"] == "failed" for result in results),
        summary_output=_display_path(config["summary_output"]),
    )
    return {
        "status": status,
        "name": config["name"],
        "task_type": config["task_type"],
        "requested_items": len(config["items"]),
        "successful_items": successful,
        "warning_items": warnings,
        "failed_items": sum(result["status"] == "failed" for result in results),
        "summary_output": _display_path(config["summary_output"]),
        "log_file": _display_path(config["log_file"]),
        "results": results,
    }


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Run a versioned MarketSignal batch task")
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--force", action="store_true", help="run even when the configured interval has not elapsed")
    parser.add_argument("--dry-run", action="store_true", help="validate the manifest without collecting data")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        summary = run_batch(args.manifest, force=args.force, dry_run=args.dry_run)
    except BatchError as exc:
        print(json.dumps({"status": "failed", "error": str(exc)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if summary["status"] in {"pass", "warning", "partial", "skipped", "validated"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
