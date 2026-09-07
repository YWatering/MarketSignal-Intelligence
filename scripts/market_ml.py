#!/usr/bin/env python3
"""Run the stage-six A-share excess-return machine-learning workflow."""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook

from marketsignal import (
    PipelineError,
    _remove_core_properties,
    _set_number_format,
    _write_table,
    clean_announcements,
    clean_financials,
    clean_news,
    clean_prices,
    fetch_online_china_adjusted_prices,
    fetch_online_china_announcements,
    fetch_online_china_financials,
    fetch_online_china_index_prices,
    fetch_online_china_news,
    filter_announcements_by_range,
    filter_financials_by_range,
    filter_news_by_range,
    filter_prices_by_range,
    parse_symbol,
)
from ml_forecasting import MLForecastError, build_excess_return_panel, run_ml_forecast_analysis
from versioning import ML_PANEL_CONTRACT_VERSION, component_versions


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CACHE_DIR = PROJECT_ROOT / ".cache" / "marketsignal"
ALLOWED_MEMBERSHIP_POLICIES = {"user_supplied_fixed_universe", "historical_constituents"}


class MLTaskError(RuntimeError):
    """Expected manifest or stage-six task failure."""


def _read_manifest(path: Path) -> dict[str, Any]:
    try:
        payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError) as exc:
        raise MLTaskError(f"could not read ML manifest: {path}") from exc
    if not isinstance(payload, dict):
        raise MLTaskError("ML manifest must contain a mapping at the top level")
    return payload


def _date_value(value: Any, field: str) -> str:
    text = str(value or "").strip()
    if text == "latest":
        return date.today().isoformat()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise MLTaskError(f"{field} must use YYYY-MM-DD or latest") from exc


def _positive_int(value: Any, field: str, minimum: int) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise MLTaskError(f"{field} must be an integer") from exc
    if number < minimum:
        raise MLTaskError(f"{field} must be at least {minimum}")
    return number


def _non_negative_float(value: Any, field: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise MLTaskError(f"{field} must be numeric") from exc
    if number < 0:
        raise MLTaskError(f"{field} cannot be negative")
    return number


def validate_manifest(payload: dict[str, Any]) -> dict[str, Any]:
    if str(payload.get("version", "")) != ML_PANEL_CONTRACT_VERSION:
        raise MLTaskError(
            f"ML manifest version must be {ML_PANEL_CONTRACT_VERSION}; received {payload.get('version', 'missing')}"
        )
    name = str(payload.get("name", "")).strip()
    if not name:
        raise MLTaskError("ML manifest requires a non-empty name")
    if str(payload.get("market", "cn_a")) != "cn_a":
        raise MLTaskError("stage-six online collection currently supports market=cn_a only")
    if str(payload.get("mode", "online")) != "online":
        raise MLTaskError("formal stage-six forecasts require mode=online")
    start_date = _date_value(payload.get("start_date"), "start_date")
    end_date = _date_value(payload.get("end_date", "latest"), "end_date")
    if start_date > end_date:
        raise MLTaskError("start_date cannot be later than end_date")
    benchmark = payload.get("benchmark")
    if not isinstance(benchmark, dict) or not str(benchmark.get("symbol", "")).strip():
        raise MLTaskError("ML manifest requires an explicit benchmark.symbol")
    benchmark_symbol = str(benchmark["symbol"]).strip()
    if not benchmark_symbol.isdigit() or len(benchmark_symbol) != 6:
        raise MLTaskError("benchmark.symbol must contain six digits")
    horizon_values = payload.get("horizons", [1, 5])
    if not isinstance(horizon_values, list) or not horizon_values:
        raise MLTaskError("horizons must be a non-empty list")
    horizons = sorted({int(value) for value in horizon_values})
    if any(value < 1 or value > 20 for value in horizons):
        raise MLTaskError("horizons must contain trading-step values from 1 to 20")
    membership_policy = str(payload.get("membership_policy", "user_supplied_fixed_universe"))
    if membership_policy not in ALLOWED_MEMBERSHIP_POLICIES:
        raise MLTaskError(
            f"membership_policy must be one of: {', '.join(sorted(ALLOWED_MEMBERSHIP_POLICIES))}"
        )
    items = payload.get("items")
    if not isinstance(items, list) or len(items) < 2:
        raise MLTaskError("ML manifest requires at least two explicitly supplied A-share items")
    if len(items) > 30:
        raise MLTaskError("ML manifest supports at most 30 items per run")
    normalized_items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            raise MLTaskError(f"item {index} must be a mapping")
        try:
            instrument = parse_symbol(str(item.get("symbol", "")), "cn_a")
        except PipelineError as exc:
            raise MLTaskError(f"item {index}: {exc}") from exc
        if instrument["symbol"] in seen:
            raise MLTaskError(f"duplicate ML panel item: {instrument['symbol']}")
        seen.add(instrument["symbol"])
        normalized_items.append(
            {
                "instrument": instrument,
                "symbol": instrument["symbol"],
                "label": str(item.get("label", instrument["symbol"])).strip() or instrument["symbol"],
                "industry": str(item.get("industry", "未分类")).strip() or "未分类",
                "size_bucket": str(item.get("size_bucket", "未分类")).strip() or "未分类",
                "membership_start": str(item.get("membership_start", "")),
                "membership_end": str(item.get("membership_end", "")),
                "membership_source": str(item.get("membership_source", membership_policy)),
            }
        )
    output = Path(str(payload.get("output", "outputs/china_ml_panel.xlsx")))
    if output.suffix.lower() != ".xlsx":
        raise MLTaskError("output must use the .xlsx extension")
    return {
        "version": ML_PANEL_CONTRACT_VERSION,
        "name": name,
        "market": "cn_a",
        "mode": "online",
        "start_date": start_date,
        "end_date": end_date,
        "benchmark": {
            "symbol": benchmark_symbol,
            "label": str(benchmark.get("label", benchmark_symbol)),
            "type": str(benchmark.get("type", "market_index")),
        },
        "horizons": horizons,
        "membership_policy": membership_policy,
        "include_context": bool(payload.get("include_context", True)),
        "final_test_dates": _positive_int(payload.get("final_test_dates", 40), "final_test_dates", 10),
        "outer_test_dates": _positive_int(payload.get("outer_test_dates", 20), "outer_test_dates", 10),
        "outer_folds": _positive_int(payload.get("outer_folds", 3), "outer_folds", 2),
        "inner_validation_dates": _positive_int(
            payload.get("inner_validation_dates", 20), "inner_validation_dates", 10
        ),
        "minimum_training_dates": _positive_int(
            payload.get("minimum_training_dates", 120), "minimum_training_dates", 60
        ),
        "transaction_cost_bps": _non_negative_float(
            payload.get("transaction_cost_bps", 10.0), "transaction_cost_bps"
        ),
        "slippage_bps": _non_negative_float(payload.get("slippage_bps", 5.0), "slippage_bps"),
        "output": output,
        "items": normalized_items,
    }


def _collect_context(
    instrument: dict[str, str],
    start_date: str,
    end_date: str,
    cache_dir: Path,
    refresh_cache: bool,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    source_rows: list[dict[str, Any]] = []
    datasets: dict[str, list[dict[str, Any]]] = {"financials": [], "news": [], "announcements": []}
    collectors = (
        ("financials", lambda: fetch_online_china_financials(instrument, cache_dir, refresh_cache)),
        ("news", lambda: fetch_online_china_news(instrument, cache_dir, refresh_cache)),
        (
            "announcements",
            lambda: fetch_online_china_announcements(
                instrument, start_date, end_date, cache_dir, refresh_cache
            ),
        ),
    )
    for dataset, collector in collectors:
        try:
            result = collector()
            raw = result[0]
            location = result[2] if dataset == "financials" else result[1]
            cache_status = result[3] if dataset == "financials" else result[2]
            provider = result[4] if dataset == "financials" else result[3]
            if dataset == "financials":
                cleaned, quality = clean_financials(raw)
                cleaned, out_of_range = filter_financials_by_range(cleaned, start_date, end_date)
            elif dataset == "news":
                cleaned, quality = clean_news(raw)
                cleaned, out_of_range = filter_news_by_range(cleaned, start_date, end_date)
            else:
                cleaned, quality = clean_announcements(raw)
                cleaned, out_of_range = filter_announcements_by_range(cleaned, start_date, end_date)
            datasets[dataset] = cleaned
            source_rows.append(
                {
                    "symbol": instrument["symbol"],
                    "dataset": dataset,
                    "provider": provider,
                    "status": "pass" if cleaned else "warning",
                    "cache_status": cache_status,
                    "source_location": location,
                    "raw_rows": quality["raw_rows"],
                    "clean_rows": quality["clean_rows"],
                    "output_rows": len(cleaned),
                    "notes": f"out_of_range={out_of_range}",
                }
            )
        except PipelineError as exc:
            source_rows.append(
                {
                    "symbol": instrument["symbol"],
                    "dataset": dataset,
                    "provider": "AKShare",
                    "status": "warning",
                    "cache_status": "error",
                    "source_location": "AKShare",
                    "raw_rows": 0,
                    "clean_rows": 0,
                    "output_rows": 0,
                    "notes": str(exc),
                }
            )
    return datasets["financials"], datasets["news"], datasets["announcements"], source_rows


def _sheet_from_dicts(workbook: Workbook, title: str, rows: list[dict[str, Any]], headers: list[str]) -> Any:
    sheet = workbook.create_sheet(title)
    _write_table(sheet, headers, [[row.get(header, "") for header in headers] for row in rows])
    return sheet


def build_ml_workbook(
    output: Path,
    config: dict[str, Any],
    result: dict[str, Any],
    source_rows: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "MarketSignal Intelligence"
    workbook.properties.title = "MarketSignal Intelligence ML excess-return report"

    readme_rows = [
        ["task_name", config["name"], "Versioned multi-stock panel task"],
        ["market", config["market"], "First implementation supports China A shares"],
        ["benchmark_symbol", config["benchmark"]["symbol"], config["benchmark"]["label"]],
        ["benchmark_type", config["benchmark"]["type"], "Explicit benchmark; no automatic substitution"],
        ["analysis_start", config["start_date"], "Business data range"],
        ["analysis_end", config["end_date"], "Business data range"],
        ["horizons", ", ".join(map(str, config["horizons"])), "Trading-step excess-return targets"],
        ["subjects", len(config["items"]), "Explicitly supplied stocks"],
        ["membership_policy", config["membership_policy"], "Controls survivorship-bias status"],
        ["pipeline_status", result["status"], "Review 数据泄漏检查 and 模型排行榜"],
        ["panel_rows", len(result["panel_rows"]), "Labeled stock-date-horizon observations"],
        ["feature_count", len(result["feature_names"]), "Point-in-time and cross-sectional features"],
        ["random_seed", result["random_seed"], "Fixed for repeatable fitted models"],
        ["limitations", result["limitations"], "Research output; not investment advice"],
    ]
    readme_rows.extend(
        [[f"version_{component}", version, "Central component version registry"] for component, version in component_versions().items()]
    )
    readme = workbook.create_sheet("README")
    _write_table(readme, ["field", "value", "notes"], readme_rows)

    panel_headers = [
        "symbol",
        "label",
        "industry",
        "size_bucket",
        "horizon",
        "feature_date",
        "target_date",
        "benchmark_symbol",
        "stock_return",
        "benchmark_return",
        "excess_return",
        "outperformed",
        "membership_source",
        "adjustment",
        *result["feature_names"],
    ]
    panel_sheet = workbook.create_sheet("面板样本")
    panel_values = []
    for row in result["panel_rows"]:
        flat = {**row, **row["features"]}
        panel_values.append([flat.get(header, "") for header in panel_headers])
    _write_table(panel_sheet, panel_headers, panel_values)
    _set_number_format(
        panel_sheet,
        {"stock_return", "benchmark_return", "excess_return", *result["feature_names"]},
        "0.000000",
    )

    forecast_headers = [
        "horizon",
        "symbol",
        "label",
        "industry",
        "size_bucket",
        "feature_date",
        "estimated_target_date",
        "benchmark_symbol",
        "model_name",
        "model_version",
        "model_status",
        "predicted_excess_return",
        "outperform_probability",
        "prediction_rank",
        "lower_bound",
        "upper_bound",
        "interval_level",
        "limitations",
    ]
    forecast_sheet = _sheet_from_dicts(workbook, "机器学习预测", result["forecasts"], forecast_headers)
    _set_number_format(
        forecast_sheet,
        {"predicted_excess_return", "outperform_probability", "lower_bound", "upper_bound"},
        "0.0000%",
    )

    leaderboard_headers = [
        "horizon",
        "model_name",
        "model_version",
        "model_family",
        "status",
        "selected",
        "locked_outer_candidate",
        "selection_reason",
        "outer_rmse",
        "outer_mae",
        "outer_oos_r2",
        "outer_ic",
        "outer_rank_ic",
        "outer_direction_accuracy",
        "outer_auc",
        "outer_brier_score",
        "outer_folds_beating_best_simple",
        "outer_positive_rank_ic_folds",
        "outer_folds",
        "final_rows",
        "final_rmse",
        "final_mae",
        "final_oos_r2",
        "final_ic",
        "final_rank_ic",
        "final_direction_accuracy",
        "final_direction_f1",
        "final_auc",
        "final_brier_score",
        "final_calibration_error",
    ]
    leaderboard_sheet = _sheet_from_dicts(workbook, "模型排行榜", result["leaderboard"], leaderboard_headers)
    _set_number_format(
        leaderboard_sheet,
        {
            "outer_rmse",
            "outer_mae",
            "outer_oos_r2",
            "outer_ic",
            "outer_rank_ic",
            "outer_direction_accuracy",
            "outer_auc",
            "outer_brier_score",
            "final_rmse",
            "final_mae",
            "final_oos_r2",
            "final_ic",
            "final_rank_ic",
            "final_direction_accuracy",
            "final_direction_f1",
            "final_auc",
            "final_brier_score",
            "final_calibration_error",
        },
        "0.000000",
    )

    rolling_headers = [
        "horizon",
        "fold",
        "model_name",
        "model_version",
        "test_start",
        "test_end",
        "parameters",
        "rows",
        "mae",
        "rmse",
        "oos_r2",
        "ic",
        "rank_ic",
        "direction_accuracy",
        "direction_f1",
        "auc",
        "brier_score",
        "calibration_error",
    ]
    rolling_sheet = _sheet_from_dicts(workbook, "滚动验证", result["rolling_validation"], rolling_headers)
    _set_number_format(
        rolling_sheet,
        {"mae", "rmse", "oos_r2", "ic", "rank_ic", "direction_accuracy", "direction_f1", "auc", "brier_score", "calibration_error"},
        "0.000000",
    )

    importance_headers = [
        "horizon",
        "model_name",
        "model_version",
        "rank",
        "feature",
        "importance",
        "signed_value",
        "interpretation",
    ]
    importance_sheet = _sheet_from_dicts(workbook, "特征重要性", result["feature_importance"], importance_headers)
    _set_number_format(importance_sheet, {"importance", "signed_value"}, "0.000000")

    explanation_headers = [
        "horizon",
        "symbol",
        "model_name",
        "rank",
        "feature",
        "feature_value",
        "local_prediction_difference",
        "direction",
        "method",
    ]
    explanation_sheet = _sheet_from_dicts(
        workbook, "预测解释", result["prediction_explanations"], explanation_headers
    )
    _set_number_format(explanation_sheet, {"feature_value", "local_prediction_difference"}, "0.000000")

    group_headers = [
        "horizon",
        "model_name",
        "model_version",
        "fold",
        "group",
        "observations",
        "average_actual_excess_return",
        "win_rate",
    ]
    group_sheet = _sheet_from_dicts(workbook, "分组检验", result["group_tests"], group_headers)
    _set_number_format(group_sheet, {"average_actual_excess_return", "win_rate"}, "0.0000%")

    cost_headers = [
        "horizon",
        "model_name",
        "model_version",
        "fold",
        "scenario",
        "periods",
        "transaction_cost_bps",
        "slippage_bps",
        "cumulative_gross_return",
        "cumulative_net_return",
        "average_spread_return",
        "win_rate",
        "turnover",
        "max_drawdown",
    ]
    cost_sheet = _sheet_from_dicts(workbook, "成本敏感性", result["cost_sensitivity"], cost_headers)
    _set_number_format(
        cost_sheet,
        {"cumulative_gross_return", "cumulative_net_return", "average_spread_return", "win_rate", "max_drawdown"},
        "0.0000%",
    )

    leakage_sheet = _sheet_from_dicts(
        workbook, "数据泄漏检查", result["leakage_checks"], ["check", "status", "details"]
    )
    final_headers = [
        "model_name",
        "model_version",
        "fold",
        "split",
        "symbol",
        "industry",
        "size_bucket",
        "feature_date",
        "target_date",
        "horizon",
        "predicted_excess_return",
        "actual_excess_return",
        "outperform_probability",
        "actual_outperformed",
        "training_rows",
        "parameters",
    ]
    final_sheet = _sheet_from_dicts(workbook, "最终测试明细", result["final_predictions"], final_headers)
    _set_number_format(
        final_sheet,
        {"predicted_excess_return", "actual_excess_return", "outperform_probability"},
        "0.0000%",
    )

    tuning_headers = [
        "horizon",
        "fold",
        "model_name",
        "parameters",
        "inner_rmse",
        "training_rows",
        "validation_rows",
        "validation_start",
        "validation_end",
    ]
    tuning_sheet = _sheet_from_dicts(workbook, "参数搜索", result["tuning_trials"], tuning_headers)
    _set_number_format(tuning_sheet, {"inner_rmse"}, "0.000000")

    source_headers = [
        "symbol",
        "dataset",
        "provider",
        "status",
        "cache_status",
        "source_location",
        "raw_rows",
        "clean_rows",
        "output_rows",
        "notes",
    ]
    _sheet_from_dicts(workbook, "数据来源", source_rows, source_headers)

    version_rows = [
        {"component": component, "version": version, "notes": "central version registry"}
        for component, version in component_versions().items()
    ]
    version_rows.extend(
        [
            {"component": "random_seed", "version": str(result["random_seed"]), "notes": "model reproducibility"},
            {
                "component": "final_test_usage",
                "version": "acceptance_only",
                "notes": result["validation_config"]["final_test_usage"],
            },
        ]
    )
    _sheet_from_dicts(workbook, "模型版本", version_rows, ["component", "version", "notes"])

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    _remove_core_properties(output)


def run_ml_task(
    manifest: Path,
    *,
    cache_dir: Path = DEFAULT_CACHE_DIR,
    refresh_cache: bool = False,
) -> dict[str, Any]:
    config = validate_manifest(_read_manifest(manifest))
    benchmark_raw, benchmark_location, benchmark_cache, benchmark_provider = fetch_online_china_index_prices(
        config["benchmark"]["symbol"],
        config["start_date"],
        config["end_date"],
        cache_dir,
        refresh_cache,
    )
    benchmark_prices, benchmark_quality = clean_prices(benchmark_raw)
    benchmark_prices, benchmark_out_of_range = filter_prices_by_range(
        benchmark_prices, config["start_date"], config["end_date"]
    )
    for row in benchmark_prices:
        row["adjustment"] = "index"
    source_rows = [
        {
            "symbol": config["benchmark"]["symbol"],
            "dataset": "benchmark_prices",
            "provider": benchmark_provider,
            "status": "pass" if benchmark_prices else "fail",
            "cache_status": benchmark_cache,
            "source_location": benchmark_location,
            "raw_rows": benchmark_quality["raw_rows"],
            "clean_rows": benchmark_quality["clean_rows"],
            "output_rows": len(benchmark_prices),
            "notes": f"out_of_range={benchmark_out_of_range}; explicit benchmark index wheels",
        }
    ]
    if not benchmark_prices:
        raise MLTaskError("benchmark source returned no clean rows in the requested range")

    subjects: list[dict[str, Any]] = []
    for item in config["items"]:
        raw_prices, price_location, price_cache, price_provider = fetch_online_china_adjusted_prices(
            item["instrument"],
            config["start_date"],
            config["end_date"],
            cache_dir,
            refresh_cache,
        )
        prices, quality = clean_prices(raw_prices)
        prices, out_of_range = filter_prices_by_range(prices, config["start_date"], config["end_date"])
        for row in prices:
            row["adjustment"] = "qfq"
        if not prices:
            raise MLTaskError(f"{item['symbol']} returned no clean adjusted price rows")
        source_rows.append(
            {
                "symbol": item["symbol"],
                "dataset": "adjusted_prices",
                "provider": price_provider,
                "status": "pass",
                "cache_status": price_cache,
                "source_location": price_location,
                "raw_rows": quality["raw_rows"],
                "clean_rows": quality["clean_rows"],
                "output_rows": len(prices),
                "notes": f"adjustment=qfq; out_of_range={out_of_range}",
            }
        )
        financials: list[dict[str, Any]] = []
        news: list[dict[str, Any]] = []
        announcements: list[dict[str, Any]] = []
        if config["include_context"]:
            financials, news, announcements, context_sources = _collect_context(
                item["instrument"],
                config["start_date"],
                config["end_date"],
                cache_dir,
                refresh_cache,
            )
            source_rows.extend(context_sources)
        subjects.append(
            {
                **{key: value for key, value in item.items() if key != "instrument"},
                "prices": prices,
                "financials": financials,
                "news": news,
                "announcements": announcements,
            }
        )

    panel_rows, latest_rows, feature_names = build_excess_return_panel(
        subjects,
        benchmark_prices,
        benchmark_symbol=config["benchmark"]["symbol"],
        benchmark_type=config["benchmark"]["type"],
        benchmark_source=benchmark_provider,
        horizons=config["horizons"],
    )
    result = run_ml_forecast_analysis(
        panel_rows,
        latest_rows,
        feature_names,
        benchmark_symbol=config["benchmark"]["symbol"],
        membership_policy=config["membership_policy"],
        final_test_dates=config["final_test_dates"],
        outer_test_dates=config["outer_test_dates"],
        outer_folds=config["outer_folds"],
        inner_validation_dates=config["inner_validation_dates"],
        minimum_training_dates=config["minimum_training_dates"],
        transaction_cost_bps=config["transaction_cost_bps"],
        slippage_bps=config["slippage_bps"],
    )
    output = config["output"]
    if not output.is_absolute():
        output = PROJECT_ROOT / output
    build_ml_workbook(output, config, result, source_rows)
    return {
        "status": result["status"],
        "output": str(output),
        "subjects": len(subjects),
        "benchmark_symbol": config["benchmark"]["symbol"],
        "panel_rows": len(panel_rows),
        "features": len(feature_names),
        "forecasts": len(result["forecasts"]),
        "selected_models": {
            str(row["horizon"]): row["model_name"]
            for row in result["leaderboard"]
            if row["selected"]
        },
        "leakage_status": {row["check"]: row["status"] for row in result["leakage_checks"]},
        "component_versions": component_versions(),
    }


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Train and validate stage-six A-share excess-return models from an explicit panel manifest"
    )
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--cache-dir", type=Path, default=DEFAULT_CACHE_DIR)
    parser.add_argument("--refresh-cache", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        if args.validate_only:
            config = validate_manifest(_read_manifest(args.manifest))
            summary = {
                "status": "valid",
                "name": config["name"],
                "subjects": len(config["items"]),
                "benchmark_symbol": config["benchmark"]["symbol"],
                "horizons": config["horizons"],
            }
        else:
            summary = run_ml_task(
                args.manifest,
                cache_dir=args.cache_dir,
                refresh_cache=args.refresh_cache,
            )
    except (MLTaskError, MLForecastError, PipelineError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
