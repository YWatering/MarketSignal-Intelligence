from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import marketsignal
from market_calendar import is_trading_day, next_trading_date
from market_ml import build_ml_workbook, validate_manifest
from ml_forecasting import build_excess_return_panel, run_ml_forecast_analysis
from tests.test_stage_six import _nonlinear_panel


class StageEightTests(unittest.TestCase):
    def test_market_manifests_normalize_and_keep_currency(self) -> None:
        cases = (
            ("cn_b", "900901.SH", "900905.SH", "USD", "CN_B_SHARE"),
            ("hk", "00700.HK", "02800.HK", "HKD", "HKEX"),
            ("us", "AAPL", "SPY", "USD", "NYSE_NASDAQ"),
        )
        for market, symbol, benchmark, currency, calendar in cases:
            with self.subTest(market=market):
                config = validate_manifest(
                    {
                        "version": "2.0",
                        "name": f"{market}-test",
                        "task_type": "single_asset",
                        "market": market,
                        "mode": "online",
                        "symbol": symbol,
                        "start_date": "2023-01-01",
                        "end_date": "latest",
                        "benchmark": {"symbol": benchmark, "type": "benchmark_asset"},
                    }
                )
                self.assertEqual(market, config["market"])
                self.assertEqual(currency, config["items"][0]["currency"])
                self.assertEqual(currency, config["benchmark"]["currency"])
                self.assertEqual(calendar, config["calendar"])
                self.assertEqual(market, config["benchmark"]["market"])

    def test_cross_market_and_cross_currency_benchmarks_are_rejected(self) -> None:
        with self.assertRaisesRegex(Exception, "benchmark.market"):
            validate_manifest(
                {
                    "version": "2.0",
                    "name": "mixed-market",
                    "task_type": "single_asset",
                    "market": "hk",
                    "mode": "online",
                    "symbol": "00700.HK",
                    "start_date": "2023-01-01",
                    "end_date": "latest",
                    "benchmark": {"symbol": "SPY", "market": "us", "type": "benchmark_asset"},
                }
            )

    def test_same_market_panel_is_allowed(self) -> None:
        config = validate_manifest(
            {
                "version": "2.0",
                "name": "hk-panel",
                "task_type": "panel",
                "market": "hk",
                "mode": "online",
                "start_date": "2023-01-01",
                "end_date": "latest",
                "benchmark": {"symbol": "02800.HK", "type": "benchmark_asset"},
                "items": [{"symbol": "00700.HK"}, {"symbol": "00941.HK"}],
            }
        )
        self.assertEqual("panel", config["task_type"])
        self.assertEqual(["00700.HK", "00941.HK"], [item["symbol"] for item in config["items"]])
        with self.assertRaisesRegex(Exception, "currency"):
            validate_manifest(
                {
                    "version": "2.0",
                    "name": "mixed-currency",
                    "task_type": "single_asset",
                    "market": "cn_b",
                    "mode": "online",
                    "symbol": "900901.SH",
                    "start_date": "2023-01-01",
                    "end_date": "latest",
                    "benchmark": {"symbol": "200002.SZ", "type": "benchmark_asset"},
                }
            )

    def test_real_exchange_calendar_handles_non_weekend_closure(self) -> None:
        self.assertTrue(is_trading_day(__import__("datetime").date(2026, 9, 8), "NYSE_NASDAQ"))
        self.assertFalse(is_trading_day(__import__("datetime").date(2026, 9, 7), "NYSE_NASDAQ"))
        self.assertEqual("2026-09-08", next_trading_date("2026-09-04", 1, "NYSE_NASDAQ"))

    def test_us_adjusted_price_route_marks_adjustment(self) -> None:
        payload = {
            "values": [
                {
                    "datetime": "2026-09-04",
                    "open": "100",
                    "high": "101",
                    "low": "99",
                    "close": "100.5",
                    "volume": "1000",
                }
            ]
        }
        with patch.object(marketsignal, "_fetch_json", return_value=(payload, "miss")) as mocked:
            rows, _, _, provider = marketsignal.fetch_online_us_adjusted_prices(
                "AAPL",
                "2026-09-01",
                "2026-09-04",
                "key",
                5,
                Path(tempfile.mkdtemp()),
                False,
            )
        self.assertEqual("Twelve Data adjusted_all", provider)
        self.assertEqual("adjusted_all", rows[0]["adjustment"])
        requested_url = mocked.call_args.args[0]
        self.assertIn("adjustment=all", requested_url)

    def test_multimarket_engine_and_workbook_keep_audit_fields(self) -> None:
        subjects, benchmark = _nonlinear_panel(count=260, stocks=1)
        subject = subjects[0]
        subject["symbol"] = "AAPL"
        subject["label"] = "Apple"
        subject["market"] = "us"
        subject["currency"] = "USD"
        for row in subject["prices"]:
            row["market"] = "us"
            row["currency"] = "USD"
            row["adjustment"] = "adjusted_all"
        for row in benchmark:
            row["market"] = "us"
            row["currency"] = "USD"
            row["adjustment"] = "adjusted_all"
        panel, latest, feature_names = build_excess_return_panel(
            [subject],
            benchmark,
            benchmark_symbol="SPY",
            benchmark_type="benchmark_asset",
            benchmark_source="deterministic benchmark",
            horizons=(1,),
            single_asset=True,
            market="us",
            calendar="NYSE_NASDAQ",
            timezone="America/New_York",
            benchmark_market="us",
            benchmark_currency="USD",
            benchmark_calendar="NYSE_NASDAQ",
            benchmark_timezone="America/New_York",
        )
        result = run_ml_forecast_analysis(
            panel,
            latest,
            feature_names,
            benchmark_symbol="SPY",
            membership_policy="single_asset",
            final_test_dates=20,
            outer_test_dates=15,
            outer_folds=2,
            inner_validation_dates=12,
            minimum_training_dates=100,
            transaction_cost_bps=0,
            slippage_bps=0,
            single_asset=True,
            market="us",
            calendar="NYSE_NASDAQ",
            benchmark_market="us",
            benchmark_currency="USD",
            benchmark_adjustment="adjusted_all",
        )
        self.assertEqual("pass", result["status"])
        self.assertTrue(all(row["market"] == "us" for row in result["forecasts"]))
        self.assertTrue(all(row["estimated_target_date"] for row in result["forecasts"]))
        self.assertTrue(all(row["status"] == "pass" for row in result["leakage_checks"]))

        config = validate_manifest(
            {
                "version": "2.0",
                "name": "us-test",
                "task_type": "single_asset",
                "market": "us",
                "mode": "online",
                "symbol": "AAPL",
                "start_date": "2023-01-01",
                "end_date": "latest",
                "benchmark": {"symbol": "SPY", "type": "benchmark_asset"},
            }
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "us.xlsx"
            build_ml_workbook(output, config, result, [])
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertIn("市场适配检查", workbook.sheetnames)
            self.assertIn("benchmark_adjustment", [cell.value for cell in workbook["单股样本"][1]])
            workbook.close()


if __name__ == "__main__":
    unittest.main()
