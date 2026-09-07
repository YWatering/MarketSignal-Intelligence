from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from market_ml import build_ml_workbook, validate_manifest
from ml_forecasting import build_excess_return_panel, run_ml_forecast_analysis


def _business_dates(count: int) -> list[str]:
    dates: list[str] = []
    current = datetime(2024, 1, 2)
    while len(dates) < count:
        if current.weekday() < 5:
            dates.append(current.strftime("%Y-%m-%d"))
        current += timedelta(days=1)
    return dates


def _price_rows(closes: list[float], volumes: list[float], dates: list[str], source: str) -> list[dict[str, object]]:
    return [
        {
            "date": trading_date,
            "open": close,
            "high": close * 1.01,
            "low": close * 0.99,
            "close": close,
            "volume": volume,
            "source": source,
            "market": "cn_a",
            "currency": "CNY",
            "adjustment": "qfq" if source != "benchmark" else "index",
        }
        for trading_date, close, volume in zip(dates, closes, volumes)
    ]


def _nonlinear_panel(count: int = 300, stocks: int = 6) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    dates = _business_dates(count)
    benchmark_close = 100.0
    benchmark_closes: list[float] = []
    benchmark_returns: list[float] = []
    for index in range(count):
        benchmark_return = 0.0004 * np.sin(index / 9.0)
        benchmark_returns.append(float(benchmark_return))
        benchmark_close *= 1.0 + benchmark_return
        benchmark_closes.append(benchmark_close)
    benchmark = _price_rows(benchmark_closes, [2_000_000.0] * count, dates, "benchmark")
    subjects: list[dict[str, object]] = []
    for stock_index in range(stocks):
        generator = np.random.default_rng(2000 + stock_index)
        volume_high = generator.integers(0, 2, size=count).astype(bool)
        close = 80.0 + stock_index * 10
        excess = 0.006 if stock_index % 2 == 0 else -0.006
        closes: list[float] = []
        volumes: list[float] = []
        for index in range(count):
            close *= 1.0 + benchmark_returns[index] + excess
            closes.append(close)
            volumes.append(2_000_000.0 if volume_high[index] else 450_000.0)
            same_sign = (excess > 0) == bool(volume_high[index])
            excess = 0.006 if same_sign else -0.006
        subjects.append(
            {
                "symbol": f"{600000 + stock_index}.SH",
                "label": f"测试股票{stock_index}",
                "industry": f"行业{stock_index % 3}",
                "size_bucket": "大市值" if stock_index < 3 else "中市值",
                "membership_source": "historical_constituents",
                "prices": _price_rows(closes, volumes, dates, f"stock-{stock_index}"),
                "financials": [],
                "news": [],
                "announcements": [],
            }
        )
    return subjects, benchmark


class StageSixTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        subjects, benchmark = _nonlinear_panel()
        panel, latest, feature_names = build_excess_return_panel(
            subjects,
            benchmark,
            benchmark_symbol="000300",
            benchmark_type="market_index",
            benchmark_source="deterministic benchmark",
            horizons=(1,),
        )
        cls.panel = panel
        cls.latest = latest
        cls.feature_names = feature_names
        cls.result = run_ml_forecast_analysis(
            panel,
            latest,
            feature_names,
            benchmark_symbol="000300",
            membership_policy="historical_constituents",
            final_test_dates=20,
            outer_test_dates=15,
            outer_folds=2,
            inner_validation_dates=12,
            minimum_training_dates=100,
            transaction_cost_bps=0,
            slippage_bps=0,
        )

    def test_manifest_requires_explicit_panel_and_benchmark(self) -> None:
        config = validate_manifest(
            {
                "version": "1.0",
                "name": "test-panel",
                "market": "cn_a",
                "mode": "online",
                "start_date": "2023-01-01",
                "end_date": "latest",
                "benchmark": {"symbol": "000300"},
                "items": [{"symbol": "600519"}, {"symbol": "000858"}],
            }
        )
        self.assertEqual("000300", config["benchmark"]["symbol"])
        self.assertEqual(2, len(config["items"]))

    def test_panel_uses_adjusted_relative_labels_and_point_in_time_keys(self) -> None:
        row = self.panel[0]
        self.assertAlmostEqual(row["stock_return"] - row["benchmark_return"], row["excess_return"])
        self.assertLess(row["feature_date"], row["target_date"])
        self.assertEqual("qfq", row["adjustment"])
        self.assertIn("rolling_beta_60d", row["features"])
        self.assertIn("rank::stock_return_20d", row["features"])
        self.assertEqual(6, len(self.latest))
        self.assertEqual(1, len({row["feature_date"] for row in self.latest}))

    def test_nonlinear_model_is_evaluated_without_forced_selection(self) -> None:
        model_names = {row["model_name"] for row in self.result["leaderboard"]}
        self.assertEqual(
            {
                "zero_excess",
                "excess_historical_mean",
                "panel_ridge",
                "elastic_net",
                "random_forest",
                "hist_gradient_boosting",
            },
            model_names,
        )
        self.assertEqual(1, sum(row["selected"] for row in self.result["leaderboard"]))
        self.assertTrue(all(row["status"] == "pass" for row in self.result["leakage_checks"]))
        nonlinear = [
            row
            for row in self.result["leaderboard"]
            if row["model_name"] in {"random_forest", "hist_gradient_boosting"}
        ]
        self.assertTrue(any(row["outer_folds_beating_best_simple"] >= 1 for row in nonlinear))

    def test_workbook_contains_auditable_stage_six_sheets_without_core_dates(self) -> None:
        config = {
            "name": "test-panel",
            "market": "cn_a",
            "benchmark": {"symbol": "000300", "label": "沪深300", "type": "market_index"},
            "start_date": "2024-01-02",
            "end_date": "2025-01-01",
            "horizons": [1],
            "items": [{"symbol": row["symbol"]} for row in self.latest],
            "membership_policy": "historical_constituents",
        }
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "stage_six.xlsx"
            build_ml_workbook(output, config, self.result, [])
            with zipfile.ZipFile(output) as archive:
                names = set(archive.namelist())
                self.assertNotIn("docProps/core.xml", names)
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertTrue(
                {
                    "机器学习预测",
                    "模型排行榜",
                    "滚动验证",
                    "特征重要性",
                    "预测解释",
                    "分组检验",
                    "成本敏感性",
                    "数据泄漏检查",
                    "模型版本",
                }.issubset(set(workbook.sheetnames))
            )
            workbook.close()


if __name__ == "__main__":
    unittest.main()
