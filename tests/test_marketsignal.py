from __future__ import annotations

import json
import gzip
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta
from unittest.mock import patch
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from marketsignal import (
    PipelineError,
    _redact_url,
    _cache_path,
    _fetch_text,
    _normalize_dataframe_news_rows,
    _normalize_dataframe_price_rows,
    _sec_headers,
    clean_financials,
    clean_news,
    clean_prices,
    extract_china_financial_rows,
    extract_financial_rows,
    parse_symbol,
    run_pipeline,
)
from forecasting import ForecastError, build_feature_vector, run_forecast_analysis


def _forecast_prices(count: int = 180, source: str = "AKShare online test") -> list[dict[str, object]]:
    start = datetime(2025, 1, 2)
    rows: list[dict[str, object]] = []
    close = 100.0
    for index in range(count):
        date = start + timedelta(days=index)
        close *= 1.0 + 0.001 + (0.004 if index % 9 == 0 else -0.002 if index % 13 == 0 else 0.0)
        rows.append(
            {
                "date": date.strftime("%Y-%m-%d"),
                "open": close * 0.998,
                "high": close * 1.006,
                "low": close * 0.994,
                "close": close,
                "volume": 1_000_000 + index * 1_000,
                "source": source,
                "market": "cn_a",
                "currency": "CNY",
            }
        )
    return rows


def _forecast_financials() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for period_end, filed_date, revenue, income, assets, liabilities, cash_flow in (
        ("2024-12-31", "2025-03-20", 1000, 200, 2000, 800, 260),
        ("2025-06-30", "2025-08-20", 600, 130, 2200, 850, 170),
    ):
        period_start = period_end[:4] + "-01-01"
        values = {
            "revenue": revenue,
            "net_income": income,
            "assets": assets,
            "liabilities": liabilities,
            "current_assets": assets * 0.5,
            "current_liabilities": liabilities * 0.5,
            "cash_and_equivalents": assets * 0.2,
            "operating_cash_flow": cash_flow,
        }
        for metric, value in values.items():
            rows.append(
                {
                    "period_start": period_start,
                    "period_end": period_end,
                    "filed_date": filed_date,
                    "fiscal_year": period_end[:4],
                    "fiscal_period": "报告期",
                    "form": "报告期",
                    "metric": metric,
                    "metric_label": metric,
                    "value": value,
                    "unit": "CNY",
                    "accession_number": "",
                    "source": "AKShare online test",
                    "market": "cn_a",
                    "currency": "CNY",
                }
            )
    return rows


class MarketSignalTests(unittest.TestCase):
    def test_clean_prices_removes_duplicate_and_invalid_rows(self) -> None:
        rows = [
            {"date": "2026-08-11", "open": "10", "high": "12", "low": "9", "close": "11", "volume": "10"},
            {"date": "2026-08-11", "open": "10", "high": "12", "low": "9", "close": "11", "volume": "10"},
            {"date": "2026-08-12", "open": "10", "high": "8", "low": "9", "close": "9", "volume": "10"},
        ]
        cleaned, quality = clean_prices(rows)
        self.assertEqual(1, len(cleaned))
        self.assertEqual(1, quality["duplicates_removed"])
        self.assertEqual(1, quality["invalid_rows"])

    def test_clean_news_removes_duplicate_and_missing_title(self) -> None:
        rows = [
            {"title": "Headline", "url": "https://example.com/1", "published_at": "time"},
            {"title": "Headline", "url": "https://example.com/1", "published_at": "time"},
            {"title": "", "url": "https://example.com/2", "published_at": "time"},
        ]
        cleaned, quality = clean_news(rows)
        self.assertEqual(1, len(cleaned))
        self.assertEqual(1, quality["duplicates_removed"])
        self.assertEqual(1, quality["invalid_rows"])

    def test_fixture_pipeline_generates_required_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "report.xlsx"
            summary = run_pipeline(
                symbol="AAPL",
                start_date="2026-08-03",
                end_date="2026-08-14",
                mode="fixture",
                output=output,
            )
            self.assertTrue(output.exists())
            self.assertEqual("pass", summary["status"])
            self.assertEqual(5, summary["price_rows"])
            self.assertEqual(2, summary["news_rows"])
            self.assertEqual(1, summary["price_quality"]["duplicates_removed"])
            self.assertEqual(1, summary["price_quality"]["invalid_rows"])
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                ["README", "主体信息", "行情数据", "财务数据", "新闻舆情", "公告数据", "指标分析", "数据来源", "数据质量", "版本信息"],
                workbook.sheetnames,
            )
            self.assertEqual("field", workbook["README"]["A1"].value)
            self.assertEqual("symbol", workbook["README"]["A2"].value)
            self.assertEqual("AAPL", workbook["README"]["B2"].value)
            self.assertEqual("Apple Inc.", workbook["主体信息"]["B2"].value)
            self.assertEqual(6, workbook["行情数据"].max_row)
            self.assertEqual(11, workbook["财务数据"].max_row)
            self.assertEqual(3, workbook["新闻舆情"].max_row)
            self.assertEqual(2, workbook["公告数据"].max_row)
            quality_rows = list(workbook["数据质量"].iter_rows(min_row=2, values_only=True))
            self.assertIn(("prices", "out_of_range", 0, "processing quality metric"), quality_rows)
            self.assertIn(("news", "out_of_range", 0, "processing quality metric"), quality_rows)
            self.assertIn(("pipeline", "status", "pass", "warning means one or more required datasets have no clean rows"), quality_rows)
            self.assertGreater(workbook["指标分析"].max_row, 1)
            self.assertEqual(6, workbook["数据来源"].max_row)
            source_locations = [
                row[5] for row in workbook["数据来源"].iter_rows(min_row=2, values_only=True)
            ]
            self.assertTrue(all(not str(location).startswith("fixture:/") for location in source_locations))
            with zipfile.ZipFile(output) as archive:
                self.assertNotIn("docProps/core.xml", archive.namelist())

    def test_fixture_pipeline_applies_requested_date_range(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            summary = run_pipeline(
                symbol="AAPL",
                start_date="2026-08-12",
                end_date="2026-08-14",
                mode="fixture",
                output=Path(temporary_directory) / "report.xlsx",
            )
            self.assertEqual(3, summary["price_rows"])
            self.assertEqual(2, summary["news_rows"])
            self.assertEqual(2, summary["price_quality"]["out_of_range"])
            self.assertEqual(0, summary["news_quality"]["out_of_range"])

    def test_fixture_data_cannot_run_formal_forecast(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            with self.assertRaisesRegex(PipelineError, "requires online mode"):
                run_pipeline(
                    symbol="AAPL",
                    start_date="2026-08-03",
                    end_date="2026-08-14",
                    mode="fixture",
                    output=Path(temporary_directory) / "report.xlsx",
                    forecast=True,
                )

    def test_invalid_range_fails_before_pipeline_execution(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            with self.assertRaises(PipelineError):
                run_pipeline(
                    symbol="AAPL",
                    start_date="2026-08-14",
                    end_date="2026-08-03",
                    mode="fixture",
                    output=Path(temporary_directory) / "report.xlsx",
                )

    def test_financial_rows_are_normalized(self) -> None:
        payload = json.loads((ROOT / "fixtures" / "aapl_financials.json").read_text(encoding="utf-8"))
        rows = extract_financial_rows(payload, "fixture")
        cleaned, quality = clean_financials(rows)
        self.assertEqual(10, len(cleaned))
        self.assertEqual(0, quality["invalid_rows"])
        self.assertTrue({row["metric"] for row in cleaned} >= {"revenue", "net_income", "assets"})

    def test_financial_rows_keep_distinct_reporting_periods(self) -> None:
        rows = [
            {
                "period_start": "2026-04-01",
                "period_end": "2026-06-30",
                "filed_date": "2026-08-01",
                "form": "10-Q",
                "metric": "revenue",
                "metric_label": "Revenue",
                "value": 94,
                "unit": "USD",
            },
            {
                "period_start": "2026-01-01",
                "period_end": "2026-06-30",
                "filed_date": "2026-08-01",
                "form": "10-Q",
                "metric": "revenue",
                "metric_label": "Revenue",
                "value": 182,
                "unit": "USD",
            },
        ]
        cleaned, quality = clean_financials(rows)
        self.assertEqual(2, len(cleaned))
        self.assertEqual(0, quality["duplicates_removed"])
        self.assertEqual({"2026-04-01", "2026-01-01"}, {row["period_start"] for row in cleaned})

    def test_market_symbol_parsing_covers_china_and_us_markets(self) -> None:
        self.assertEqual(
            {"symbol": "600519.SH", "market": "cn_a", "currency": "CNY", "exchange": "SSE"},
            {key: parse_symbol("600519")[key] for key in ("symbol", "market", "currency", "exchange")},
        )
        self.assertEqual(
            {"symbol": "900901.SH", "market": "cn_b", "currency": "USD", "exchange": "SSE"},
            {key: parse_symbol("900901")[key] for key in ("symbol", "market", "currency", "exchange")},
        )
        self.assertEqual(
            {"symbol": "200002.SZ", "market": "cn_b", "currency": "HKD", "exchange": "SZSE"},
            {key: parse_symbol("200002")[key] for key in ("symbol", "market", "currency", "exchange")},
        )
        self.assertEqual(
            {"symbol": "00700.HK", "market": "hk", "currency": "HKD", "exchange": "HKEX"},
            {key: parse_symbol("HK00700")[key] for key in ("symbol", "market", "currency", "exchange")},
        )
        self.assertEqual("us", parse_symbol("NASDAQ:AAPL")["market"])

    def test_market_symbol_parsing_rejects_conflicting_exchange(self) -> None:
        with self.assertRaises(PipelineError):
            parse_symbol("600519.HK")
        with self.assertRaises(PipelineError):
            parse_symbol("200002.SH")
        with self.assertRaises(PipelineError):
            parse_symbol("600519", market="cn_b")

    def test_china_price_and_news_rows_are_normalized(self) -> None:
        prices = _normalize_dataframe_price_rows(
            [
                {"日期": "2026-08-12", "开盘": "1,000", "最高": 1010, "最低": 990, "收盘": "1005", "成交量": 1234},
            ],
            "AKShare stock_zh_a_hist_tx",
            "cn_a",
            "CNY",
        )
        self.assertEqual("2026-08-12", prices[0]["date"])
        self.assertEqual(1000.0, prices[0]["open"])
        self.assertEqual("CNY", prices[0]["currency"])
        news = _normalize_dataframe_news_rows(
            [
                {
                    "新闻标题": "贵州茅台营收增长，市场看好",
                    "新闻内容": "公司盈利改善并创纪录",
                    "发布时间": "2026-08-12 09:30:00",
                    "文章来源": "测试来源",
                    "新闻链接": "https://example.test/maotai",
                }
            ],
            "AKShare stock_news_em",
            "cn_a",
            "CNY",
        )
        cleaned, quality = clean_news(news)
        self.assertEqual(1, len(cleaned))
        self.assertEqual("positive", cleaned[0]["sentiment"])
        self.assertIn("增长", cleaned[0]["sentiment_basis"])
        self.assertEqual(0, quality["invalid_rows"])

    def test_china_financial_tables_map_to_common_metrics(self) -> None:
        instrument = parse_symbol("600519")
        tables = {
            "profit": [
                {
                    "REPORT_DATE": "2026-06-30",
                    "NOTICE_DATE": "2026-08-29",
                    "REPORT_TYPE": "中报",
                    "TOTAL_OPERATE_INCOME": "90,000,000,000",
                    "PARENT_NETPROFIT": 45_000_000_000,
                    "CURRENCY": "CNY",
                }
            ],
            "balance": [
                {
                    "REPORT_DATE": "2026-06-30",
                    "NOTICE_DATE": "2026-08-29",
                    "REPORT_TYPE": "中报",
                    "TOTAL_ASSETS": 300_000_000_000,
                    "TOTAL_LIABILITIES": 100_000_000_000,
                    "TOTAL_CURRENT_ASSETS": 150_000_000_000,
                    "TOTAL_CURRENT_LIAB": 80_000_000_000,
                    "MONETARYFUNDS": 50_000_000_000,
                    "CURRENCY": "CNY",
                }
            ],
            "cash": [
                {
                    "REPORT_DATE": "2026-06-30",
                    "NOTICE_DATE": "2026-08-29",
                    "REPORT_TYPE": "中报",
                    "NETCASH_OPERATE": 40_000_000_000,
                    "CURRENCY": "CNY",
                }
            ],
        }
        rows = extract_china_financial_rows(tables, instrument)
        cleaned, quality = clean_financials(rows)
        self.assertEqual(8, len(cleaned))
        self.assertEqual(0, quality["invalid_rows"])
        self.assertEqual(
            {"revenue", "net_income", "assets", "liabilities", "current_assets", "current_liabilities", "cash_and_equivalents", "operating_cash_flow"},
            {row["metric"] for row in cleaned},
        )
        self.assertTrue(all(row["market"] == "cn_a" and row["currency"] == "CNY" for row in cleaned))

    def test_hong_kong_annual_report_code_is_normalized(self) -> None:
        instrument = parse_symbol("00700.HK")
        rows = extract_china_financial_rows(
            {
                "profit": [
                    {
                        "REPORT_DATE": "2025-12-31 00:00:00",
                        "DATE_TYPE_CODE": "001",
                        "START_DATE": "2025-01-01 00:00:00",
                        "STD_ITEM_NAME": "营业额",
                        "AMOUNT": 100,
                    }
                ]
            },
            instrument,
        )
        cleaned, quality = clean_financials(rows)
        self.assertEqual(1, len(cleaned))
        self.assertEqual("年度", cleaned[0]["form"])
        self.assertEqual(0, quality["invalid_rows"])

    def test_forecast_features_respect_financial_filing_date(self) -> None:
        prices = _forecast_prices(90)
        financials = _forecast_financials()
        before_filing = build_feature_vector(
            [row for row in prices if str(row["date"]) <= "2025-03-10"],
            financials,
            [],
            [],
        )
        after_filing = build_feature_vector(
            [row for row in prices if str(row["date"]) <= "2025-04-10"],
            financials,
            [],
            [],
        )
        self.assertEqual(0.0, before_filing["fundamental_coverage"])
        self.assertGreater(after_filing["fundamental_coverage"], 0.0)
        self.assertAlmostEqual(0.2, after_filing["net_profit_margin"])

    def test_forecast_analysis_produces_models_backtest_and_intervals(self) -> None:
        prices = _forecast_prices()
        result = run_forecast_analysis(
            prices,
            _forecast_financials(),
            [
                {
                    "published_at": "2025-05-20 09:00:00",
                    "sentiment": "positive",
                    "source": "AKShare online test",
                }
            ],
            [
                {
                    "filed_date": "2025-05-21",
                    "source": "AKShare online test",
                }
            ],
            horizon=3,
            minimum_history=120,
            validation_points=20,
            ridge_alpha=1.0,
            market="cn_a",
            currency="CNY",
        )
        self.assertEqual("pass", result["status"])
        self.assertEqual(6, len(result["forecasts"]))
        self.assertEqual(2, len(result["evaluations"]))
        self.assertEqual(40, len(result["backtest"]))
        self.assertEqual(14, len(result["feature_contributions"]))
        self.assertIn(result["selected_model"], {"last_close_baseline", "multisignal_ridge"})
        self.assertTrue(all(row["lower_bound"] <= row["predicted_close"] <= row["upper_bound"] for row in result["forecasts"]))

    def test_forecast_analysis_rejects_fixture_price_sources(self) -> None:
        with self.assertRaisesRegex(ForecastError, "cannot use fixture"):
            run_forecast_analysis(
                _forecast_prices(source="fixture:prices"),
                [],
                [],
                [],
                minimum_history=120,
                market="cn_a",
                currency="CNY",
            )

    def test_forecast_analysis_rejects_insufficient_history(self) -> None:
        with self.assertRaisesRegex(ForecastError, "requires at least 120"):
            run_forecast_analysis(
                _forecast_prices(count=100),
                _forecast_financials(),
                [],
                [],
                minimum_history=120,
                market="cn_a",
                currency="CNY",
            )

    def test_forecast_analysis_honors_requested_validation_window(self) -> None:
        result = run_forecast_analysis(
            _forecast_prices(count=100),
            _forecast_financials(),
            [],
            [],
            minimum_history=80,
            validation_points=30,
            market="cn_a",
            currency="CNY",
        )
        self.assertEqual(30, result["validation_samples"])
        self.assertEqual(60, len(result["backtest"]))

    def test_online_pipeline_writes_stage_three_workbook(self) -> None:
        prices = _forecast_prices()
        entity = {
            "symbol": "600519.SH",
            "provider_symbol": "600519",
            "market": "cn_a",
            "market_name": "中国A股",
            "company_name": "贵州茅台",
            "cik": "",
            "exchange": "SSE",
            "sic": "",
            "sic_description": "",
            "currency": "CNY",
            "source": "AKShare",
        }
        with tempfile.TemporaryDirectory() as temporary_directory, patch(
            "marketsignal.fetch_online_china_prices",
            return_value=(prices, "https://example.test/prices", "miss", "AKShare online test"),
        ), patch(
            "marketsignal.fetch_online_china_news",
            return_value=([], "https://example.test/news", "miss", "AKShare online test"),
        ), patch(
            "marketsignal.fetch_online_china_financials",
            return_value=(
                _forecast_financials(),
                entity,
                "https://example.test/financials",
                "miss",
                "AKShare online test",
                "test financials",
            ),
        ), patch(
            "marketsignal.fetch_online_china_announcements",
            return_value=([], "https://example.test/notices", "not_supported", "test notice adapter"),
        ):
            output = Path(temporary_directory) / "forecast.xlsx"
            summary = run_pipeline(
                symbol="600519",
                start_date=None,
                end_date=None,
                mode="online",
                output=output,
                forecast=True,
                forecast_horizon=3,
                forecast_minimum_history=120,
                forecast_validation_points=20,
            )
            self.assertEqual("pass", summary["forecast_status"])
            self.assertEqual(6, summary["forecast_rows"])
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(
                [
                    "README",
                    "主体信息",
                    "行情数据",
                    "财务数据",
                    "新闻舆情",
                    "公告数据",
                    "指标分析",
                    "预测结果",
                    "模型评估",
                    "回测明细",
                    "特征贡献",
                    "数据来源",
                    "数据质量",
                    "版本信息",
                ],
                workbook.sheetnames,
            )
            self.assertGreater(workbook["预测结果"].max_row, 1)
            self.assertEqual(3, workbook["模型评估"].max_row)
            self.assertGreater(workbook["回测明细"].max_row, 1)
            self.assertEqual(15, workbook["特征贡献"].max_row)

    def test_sec_user_agent_and_url_redaction(self) -> None:
        self.assertEqual(
            "https://example.test/data?apikey=redacted&symbol=AAPL",
            _redact_url("https://example.test/data?apikey=secret&symbol=AAPL"),
        )
        self.assertEqual(
            {"User-Agent": "MarketSignal test@example.com", "Accept-Encoding": "gzip, deflate"},
            _sec_headers("MarketSignal test@example.com"),
        )
        with self.assertRaises(PipelineError):
            _sec_headers("MarketSignal without contact")

    def test_response_cache_avoids_repeat_request(self) -> None:
        class FakeResponse:
            class Headers:
                def get(self, name, default=""):
                    return "gzip" if name == "Content-Encoding" else default

            headers = Headers()

            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc_value, traceback):
                return False

            def read(self):
                return gzip.compress(b"{\"ok\": true}")

        with tempfile.TemporaryDirectory() as temporary_directory:
            cache_dir = Path(temporary_directory)
            url = "https://example.test/payload"
            with patch("marketsignal.urllib.request.urlopen", return_value=FakeResponse()) as request:
                content, cache_status = _fetch_text(url, 1, cache_dir=cache_dir)
                self.assertEqual("miss", cache_status)
                self.assertEqual('{"ok": true}', content)
                self.assertEqual(1, request.call_count)
            content, cache_status = _fetch_text(url, 1, cache_dir=cache_dir)
            self.assertEqual("hit", cache_status)
            self.assertEqual('{"ok": true}', content)
            self.assertTrue(_cache_path(cache_dir, url).exists())


if __name__ == "__main__":
    unittest.main()
