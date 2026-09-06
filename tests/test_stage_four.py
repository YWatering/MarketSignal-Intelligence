from __future__ import annotations

import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from market_batch import BatchError, run_batch, validate_manifest
from marketsignal import PipelineError, run_pipeline
from operations import compare_price_sources, schedule_is_due
from versioning import BATCH_CONTRACT_VERSION, component_versions


def _manifest_text(root: Path, *, items: str, mode: str = "online", schedule: bool = False) -> str:
    lines = [
        f'version: "{BATCH_CONTRACT_VERSION}"',
        "name: stage-four-test",
        "task_type: industry",
        "group: test-group",
        f"mode: {mode}",
        f'output_dir: "{root / "items"}"',
        f'summary_output: "{root / "summary.xlsx"}"',
        f'log_file: "{root / "run.jsonl"}"',
        f'state_file: "{root / "state.json"}"',
        "cross_validate_prices: false",
        "execution:",
        "  attempts: 2",
        "  retry_delay_seconds: 0",
        "  continue_on_error: true",
    ]
    if schedule:
        lines.extend(["schedule:", "  interval_hours: 24"])
    lines.extend(["items:", items])
    return "\n".join(lines)


def _summary(symbol: str, output: Path) -> dict[str, object]:
    workbook = Workbook()
    workbook.save(output)
    return {
        "symbol": symbol,
        "currency": "CNY",
        "latest_price_date": "2026-09-04",
        "latest_close": 100.0,
        "period_return_pct": 5.0,
        "close_return_volatility_pct": 1.2,
        "net_profit_margin_pct": 20.0,
        "current_ratio": 2.0,
        "operating_cash_flow_margin_pct": 22.0,
        "positive_news_count": 2,
        "negative_news_count": 1,
        "news_tone_balance": 1,
        "price_rows": 200,
        "financial_rows": 20,
        "news_rows": 5,
        "announcement_rows": 8,
        "forecast_status": "not_requested",
        "selected_model": "",
        "selected_forecast_close": None,
        "forecast_lower_bound": None,
        "forecast_upper_bound": None,
        "cross_validation_status": "not_requested",
        "output": str(output),
        "status": "pass",
    }


class StageFourTests(unittest.TestCase):
    def test_price_cross_validation_reports_pass_and_warning(self) -> None:
        primary = [
            {"date": "2026-09-01", "close": 100},
            {"date": "2026-09-02", "close": 101},
            {"date": "2026-09-03", "close": 102},
        ]
        secondary = [
            {"date": "2026-09-01", "close": 100.1},
            {"date": "2026-09-02", "close": 101.1},
            {"date": "2026-09-03", "close": 102.1},
        ]
        result = compare_price_sources(
            primary,
            secondary,
            primary_provider="primary",
            secondary_provider="secondary",
            tolerance_pct=1.0,
        )
        self.assertEqual("pass", result["status"])
        secondary[-1]["close"] = 90
        warning = compare_price_sources(
            primary,
            secondary,
            primary_provider="primary",
            secondary_provider="secondary",
            tolerance_pct=1.0,
        )
        self.assertEqual("warning", warning["status"])
        self.assertEqual(1, warning["outside_tolerance_rows"])

    def test_schedule_due_logic(self) -> None:
        self.assertTrue(schedule_is_due(interval_hours=24, last_completion_epoch=None, now_epoch=1000))
        self.assertFalse(schedule_is_due(interval_hours=24, last_completion_epoch=1000, now_epoch=2000))
        self.assertTrue(schedule_is_due(interval_hours=24, last_completion_epoch=1000, now_epoch=87400))

    def test_manifest_rejects_duplicate_subjects(self) -> None:
        with self.assertRaisesRegex(BatchError, "duplicate batch item"):
            validate_manifest(
                {
                    "version": BATCH_CONTRACT_VERSION,
                    "name": "duplicate-test",
                    "task_type": "theme",
                    "items": [
                        {"symbol": "600519", "market": "cn_a"},
                        {"symbol": "600519.SH", "market": "cn_a"},
                    ],
                }
            )

    def test_manifest_rejects_invalid_model_settings(self) -> None:
        with self.assertRaisesRegex(BatchError, "forecast_horizon"):
            validate_manifest(
                {
                    "version": BATCH_CONTRACT_VERSION,
                    "name": "invalid-settings",
                    "task_type": "portfolio",
                    "forecast_horizon": 0,
                    "items": [{"symbol": "600519", "market": "cn_a"}],
                }
            )

    def test_batch_runs_multiple_subjects_and_writes_summary(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            manifest = root / "batch.yaml"
            manifest.write_text(
                _manifest_text(
                    root,
                    items="  - symbol: \"600519\"\n    market: cn_a\n    label: first\n  - symbol: \"000858\"\n    market: cn_a\n    label: second",
                ),
                encoding="utf-8",
            )

            def runner(**kwargs):
                return _summary(kwargs["symbol"], kwargs["output"])

            summary = run_batch(manifest, force=True, runner=runner, now_epoch=1000)
            self.assertEqual("pass", summary["status"])
            self.assertEqual(2, summary["successful_items"])
            workbook = load_workbook(summary["summary_output"], read_only=True, data_only=True)
            self.assertEqual(["README", "主体任务", "横向比较", "版本信息"], workbook.sheetnames)
            self.assertEqual(3, workbook["主体任务"].max_row)
            self.assertEqual(3, workbook["横向比较"].max_row)
            comparison_headers = [cell.value for cell in workbook["横向比较"][1]]
            self.assertIn("period_return_pct", comparison_headers)
            self.assertIn("net_profit_margin_pct", comparison_headers)
            workbook.close()
            with zipfile.ZipFile(summary["summary_output"]) as archive:
                self.assertNotIn("docProps/core.xml", archive.namelist())
            log_lines = (root / "run.jsonl").read_text(encoding="utf-8").splitlines()
            self.assertGreaterEqual(len(log_lines), 6)
            self.assertTrue(all("run_id" in json.loads(line) for line in log_lines))

    def test_batch_retries_then_succeeds_and_schedule_skips(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            manifest = root / "batch.yaml"
            manifest.write_text(
                _manifest_text(
                    root,
                    items="  - symbol: \"600519\"\n    market: cn_a\n    label: retry",
                    schedule=True,
                ),
                encoding="utf-8",
            )
            attempts = 0

            def runner(**kwargs):
                nonlocal attempts
                attempts += 1
                if attempts == 1:
                    raise PipelineError("transient test failure")
                return _summary(kwargs["symbol"], kwargs["output"])

            first = run_batch(manifest, force=True, runner=runner, now_epoch=1000)
            self.assertEqual("pass", first["status"])
            self.assertEqual(2, first["results"][0]["attempts"])
            second = run_batch(manifest, runner=runner, now_epoch=2000)
            self.assertEqual("skipped", second["status"])
            self.assertEqual(2, attempts)

    def test_batch_isolates_failed_subject(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            manifest = root / "batch.yaml"
            manifest.write_text(
                _manifest_text(
                    root,
                    items="  - symbol: \"600519\"\n    market: cn_a\n    label: failed\n  - symbol: \"000858\"\n    market: cn_a\n    label: passed",
                ),
                encoding="utf-8",
            )

            def runner(**kwargs):
                if kwargs["symbol"] == "600519":
                    raise PipelineError("permanent test failure")
                return _summary(kwargs["symbol"], kwargs["output"])

            summary = run_batch(manifest, force=True, runner=runner, now_epoch=1000)
            self.assertEqual("partial", summary["status"])
            self.assertEqual(1, summary["successful_items"])
            self.assertEqual(1, summary["failed_items"])

    def test_batch_preserves_completed_warning_status(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            manifest = root / "batch.yaml"
            manifest.write_text(
                _manifest_text(
                    root,
                    items="  - symbol: \"600519\"\n    market: cn_a\n    label: warning",
                ),
                encoding="utf-8",
            )

            def runner(**kwargs):
                summary = _summary(kwargs["symbol"], kwargs["output"])
                summary["status"] = "warning"
                return summary

            summary = run_batch(manifest, force=True, runner=runner, now_epoch=1000)
            self.assertEqual("warning", summary["status"])
            self.assertEqual(1, summary["warning_items"])
            self.assertEqual("warning", summary["results"][0]["status"])

    def test_online_pipeline_writes_price_cross_validation(self) -> None:
        prices = [
            {
                "date": f"2026-09-0{index}",
                "open": 100 + index,
                "high": 101 + index,
                "low": 99 + index,
                "close": 100 + index,
                "volume": 1000,
                "source": "AKShare primary",
                "market": "cn_a",
                "currency": "CNY",
            }
            for index in range(1, 4)
        ]
        secondary = [{**row, "source": "AKShare secondary", "close": row["close"] + 0.1} for row in prices]
        entity = {
            "symbol": "600519.SH",
            "provider_symbol": "600519",
            "market": "cn_a",
            "market_name": "中国A股",
            "company_name": "test",
            "cik": "",
            "exchange": "SSE",
            "sic": "",
            "sic_description": "",
            "currency": "CNY",
            "source": "AKShare",
        }
        with tempfile.TemporaryDirectory() as temporary_directory, patch(
            "marketsignal.fetch_online_china_prices",
            return_value=(prices, "https://example.test/prices", "miss", "AKShare primary"),
        ), patch(
            "marketsignal.fetch_online_china_secondary_prices",
            return_value=(secondary, "https://example.test/secondary", "miss", "AKShare secondary"),
        ), patch(
            "marketsignal.fetch_online_china_news",
            return_value=([{"published_at": "2026-09-03", "title": "test", "url": "https://example.test", "summary": "", "sentiment": "neutral", "sentiment_basis": "", "source": "AKShare", "market": "cn_a", "currency": "CNY"}], "https://example.test/news", "miss", "AKShare"),
        ), patch(
            "marketsignal.fetch_online_china_financials",
            return_value=([{"period_start": "2026-01-01", "period_end": "2026-06-30", "filed_date": "2026-08-01", "fiscal_year": "2026", "fiscal_period": "中报", "form": "中报", "metric": "revenue", "metric_label": "营业收入", "value": 100, "unit": "CNY", "accession_number": "", "source": "AKShare", "market": "cn_a", "currency": "CNY"}], entity, "https://example.test/financials", "miss", "AKShare", "test"),
        ), patch(
            "marketsignal.fetch_online_china_announcements",
            return_value=([], "https://example.test/notices", "not_supported", "test"),
        ):
            output = Path(temporary_directory) / "crosscheck.xlsx"
            summary = run_pipeline(
                symbol="600519",
                market="cn_a",
                start_date=None,
                end_date=None,
                mode="online",
                output=output,
                cross_validate_prices=True,
            )
            self.assertEqual("pass", summary["cross_validation_status"])
            workbook = load_workbook(output, read_only=True, data_only=True)
            self.assertIn("交叉校验", workbook.sheetnames)
            self.assertEqual("pass", workbook["交叉校验"]["A2"].value)
            workbook.close()

    def test_component_versions_are_explicit(self) -> None:
        versions = component_versions()
        self.assertEqual(BATCH_CONTRACT_VERSION, versions["batch_contract"])
        self.assertTrue(all(value for value in versions.values()))


if __name__ == "__main__":
    unittest.main()
