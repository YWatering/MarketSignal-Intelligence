from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT))

from market_ml import build_ml_workbook, validate_manifest
from ml_forecasting import build_excess_return_panel, run_ml_forecast_analysis
from tests.test_stage_six import _nonlinear_panel


class StageSevenTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        subjects, benchmark = _nonlinear_panel()
        single_subject = [subjects[0]]
        panel, latest, feature_names = build_excess_return_panel(
            single_subject,
            benchmark,
            benchmark_symbol="000300",
            benchmark_type="market_index",
            benchmark_source="deterministic benchmark",
            horizons=(1, 5),
            single_asset=True,
        )
        cls.panel = panel
        cls.latest = latest
        cls.feature_names = feature_names
        cls.result = run_ml_forecast_analysis(
            panel,
            latest,
            feature_names,
            benchmark_symbol="000300",
            membership_policy="single_asset",
            final_test_dates=20,
            outer_test_dates=15,
            outer_folds=2,
            inner_validation_dates=12,
            minimum_training_dates=100,
            transaction_cost_bps=0,
            slippage_bps=0,
            single_asset=True,
        )

    def test_single_asset_manifest_requires_one_symbol_and_explicit_benchmark(self) -> None:
        config = validate_manifest(
            {
                "version": "1.0",
                "name": "single-test",
                "task_type": "single_asset",
                "market": "cn_a",
                "mode": "online",
                "symbol": "600519",
                "start_date": "2023-01-01",
                "end_date": "latest",
                "benchmark": {"symbol": "000300"},
                "horizons": [1, 5],
            }
        )
        self.assertEqual("single_asset", config["task_type"])
        self.assertEqual("single_asset", config["membership_policy"])
        self.assertEqual(1, len(config["items"]))
        self.assertEqual("600519.SH", config["items"][0]["symbol"])

    def test_single_asset_panel_has_no_cross_sectional_features(self) -> None:
        self.assertTrue(self.panel)
        self.assertTrue(self.latest)
        self.assertTrue(all(row["feature_date"] < row["target_date"] for row in self.panel))
        self.assertTrue(all(not name.startswith("rank::") for name in self.feature_names))
        self.assertEqual(2, len(self.latest))
        self.assertEqual(1, len({row["symbol"] for row in self.latest}))
        self.assertTrue(all(row["adjustment"] == "qfq" for row in self.panel))

    def test_single_asset_uses_relative_strategy_and_no_group_tests(self) -> None:
        self.assertEqual("single_asset", self.result["task_type"])
        self.assertEqual("pass", self.result["status"])
        self.assertEqual([], self.result["group_tests"])
        self.assertTrue(self.result["cost_sensitivity"])
        self.assertTrue(
            all(row["strategy_type"] == "single_asset_relative_sign" for row in self.result["cost_sensitivity"])
        )
        self.assertTrue(all(row["prediction_rank"] == "not_applicable" for row in self.result["forecasts"]))

    def test_single_asset_workbook_uses_dedicated_sheets(self) -> None:
        config = validate_manifest(
            {
                "version": "1.0",
                "name": "single-test",
                "task_type": "single_asset",
                "market": "cn_a",
                "mode": "online",
                "symbol": "600519",
                "start_date": "2023-01-01",
                "end_date": "latest",
                "benchmark": {"symbol": "000300"},
                "horizons": [1, 5],
            }
        )
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "single_asset.xlsx"
            config["output"] = output
            build_ml_workbook(output, config, self.result, [])
            with zipfile.ZipFile(output) as archive:
                self.assertNotIn("docProps/core.xml", archive.namelist())
            from openpyxl import load_workbook

            workbook = load_workbook(output, read_only=True, data_only=False)
            self.assertIn("单股样本", workbook.sheetnames)
            self.assertIn("单股超额收益预测", workbook.sheetnames)
            self.assertNotIn("分组检验", workbook.sheetnames)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
